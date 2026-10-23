"""Parsers: text, page numbers and sections come out right; bad files fail with the right code."""
import pytest

from app.core.config import settings
from app.ingest.errors import IngestError
from app.ingest.parsers import parse, parse_txt
from app.scripts import synthetic_docs as sd


def _sample(name):
    return next(s for s in sd.SAMPLES if s.filename == name)


def test_txt_paragraphs_and_nul_removal():
    blocks = parse("txt", b"First paragraph\nstill first.\n\n\nSecond   paragraph.")
    assert [b.text for b in blocks] == ["First paragraph still first.", "Second paragraph."]
    assert parse_txt(b"a\x00b")[0].text == "ab"  # Postgres cannot store NUL


def test_markdown_headings_become_sections():
    blocks = parse("md", sd.build(_sample("leave-policy.md")))
    sections = {b.section for b in blocks}
    assert {"Annual leave", "Sick leave", "How to apply"} <= sections
    annual = [b for b in blocks if b.section == "Annual leave" and "24 days" in b.text]
    assert annual


def test_markdown_code_fences_do_not_create_headings():
    blocks = parse("md", b"# Title\n\n```\n# not a heading\n```\n\nText")
    assert {b.section for b in blocks} == {"Title"}


def test_docx_headings_and_tables():
    blocks = parse("docx", sd.build(_sample("pricing-playbook.docx")))
    assert any(b.section == "Discount rules" and "10 percent" in b.text for b in blocks)
    import io

    from docx import Document

    doc = Document()
    doc.add_heading("Plan", level=1)
    t = doc.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Band", "Range"
    t.rows[1].cells[0].text, t.rows[1].cells[1].text = "L1", "4 to 6"
    buf = io.BytesIO()
    doc.save(buf)
    texts = [b.text for b in parse("docx", buf.getvalue())]
    assert "Band | Range" in texts and "L1 | 4 to 6" in texts


def test_pdf_keeps_page_numbers():
    sample = _sample("fy2026-budget-forecast.pdf")
    blocks = parse("pdf", sd.build(sample))
    pages = {b.page for b in blocks}
    assert pages == {1, 2, 3, 4}  # Summary, Costs, Risks and the Marker section: one per page
    marker_pages = {b.page for b in blocks if sample.canary in b.text}
    assert marker_pages == {4}  # the canary sits on the last page, and the page number says so


def test_pdf_without_text_has_no_text_error():
    with pytest.raises(IngestError) as exc:
        parse("pdf", sd.scanned_like_pdf())
    assert exc.value.code == "NO_TEXT"


def test_corrupt_pdf_and_fake_docx_are_parse_errors():
    for kind, data in (("pdf", sd.corrupt_pdf()), ("docx", sd.fake_docx())):
        with pytest.raises(IngestError) as exc:
            parse(kind, data)
        assert exc.value.code == "PARSE_ERROR"


def test_pdf_with_too_many_pages_is_refused(monkeypatch):
    monkeypatch.setattr(settings, "max_pdf_pages", 2)
    with pytest.raises(IngestError) as exc:
        parse("pdf", sd.build(_sample("fy2026-budget-forecast.pdf")))
    assert exc.value.code == "TOO_LARGE"


def test_text_with_only_whitespace_has_no_text_error():
    with pytest.raises(IngestError) as exc:
        parse("txt", b"   \n\n  \t ")
    assert exc.value.code == "NO_TEXT"


def test_error_messages_never_contain_file_content():
    secret = b"CANARY-SECRET-TEXT-123"
    with pytest.raises(IngestError) as exc:
        parse("pdf", b"%PDF-1.4 " + secret)
    assert "CANARY" not in exc.value.message and "CANARY" not in str(exc.value)
